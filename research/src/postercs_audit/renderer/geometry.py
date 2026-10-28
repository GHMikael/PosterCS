#!/usr/bin/env python3
"""Compare picture display/readability proxies in prototype before/after PPTX files."""

from __future__ import annotations

import csv
import json
import argparse
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from postercs_audit.paths import RUNS_ROOT
from postercs_audit.renderer.visual_metrics import EMU_PER_INCH, TARGET_HEIGHT, TARGET_WIDTH, _visual_metrics


OUTPUT = RUNS_ROOT / "renderer"


def summarize(path: Path) -> dict[str, Any]:
    prs = Presentation(path)
    scale_x = TARGET_WIDTH / (prs.slide_width / EMU_PER_INCH)
    scale_y = TARGET_HEIGHT / (prs.slide_height / EMU_PER_INCH)
    rows = []
    for shape in prs.slides[0].shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        width = shape.width / EMU_PER_INCH * scale_x
        height = shape.height / EMU_PER_INCH * scale_y
        visual = _visual_metrics(shape.image.blob, width, height)
        rows.append({
            "short_side": min(width, height),
            "area": width * height,
            "detail": visual["median_detail_height_1080px"],
        })
    details = [float(row["detail"]) for row in rows if row["detail"] != ""]
    return {
        "picture_count": len(rows),
        "min_short_side": round(min((r["short_side"] for r in rows), default=0), 1),
        "median_short_side": round(sorted(r["short_side"] for r in rows)[len(rows) // 2], 1) if rows else 0,
        "min_area": round(min((r["area"] for r in rows), default=0), 1),
        "min_detail": round(min(details), 2) if details else "",
        "median_detail": round(sorted(details)[len(details) // 2], 2) if details else "",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT)
    args = parser.parse_args()
    output = args.output_dir
    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    rows = []
    for decision in manifest:
        folder = decision["run_folder"]
        before_path = ROOT / "outputs" / "runs" / folder / "pptx" / "iter_1.pptx"
        after_path = output / folder / "after.pptx"
        before = summarize(before_path)
        after = summarize(after_path)
        rows.append({
            "run_folder": folder,
            "original_template": decision["original_template"],
            "render_template": decision["render_template"],
            **{f"before_{key}": value for key, value in before.items()},
            **{f"after_{key}": value for key, value in after.items()},
            "min_short_side_ratio": round(after["min_short_side"] / before["min_short_side"], 2)
            if before["min_short_side"] else "",
            "min_area_ratio": round(after["min_area"] / before["min_area"], 2) if before["min_area"] else "",
            "min_detail_ratio": round(float(after["min_detail"]) / float(before["min_detail"]), 2)
            if before["min_detail"] not in ("", 0) and after["min_detail"] != "" else "",
        })
    path = output / "geometry_comparison.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
