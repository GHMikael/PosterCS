#!/usr/bin/env python3
"""Measure figure readability proxies in 16:9 PosterCS PPTX outputs.

This is a diagnostic/calibration tool, not a ground-truth labeler. It reads the
embedded picture pixels and their final PPT display boxes, converts geometry to
a 1920x1080-equivalent canvas, and reports raw size/detail metrics. The
connected-component metric is only a proxy for small internal labels because
no OCR engine is required.
"""

from __future__ import annotations

import argparse
import csv
import io
import math
from pathlib import Path
from typing import Any

import cv2
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


from postercs_audit.paths import DATA_ROOT, RESULTS_ROOT

ROOT = DATA_ROOT
DEFAULT_LABELS = DATA_ROOT / "renderer-development" / "posters.csv"
DEFAULT_OUT = RESULTS_ROOT / "renderer-readability"
EMU_PER_INCH = 914400
TARGET_WIDTH = 1920
TARGET_HEIGHT = 1080


def _round(value: float | None, digits: int = 3) -> float | str:
    return "" if value is None or not math.isfinite(value) else round(value, digits)


def _entropy(gray: np.ndarray) -> float:
    hist = cv2.calcHist([gray], [0], None, [256], [0, 256]).ravel()
    probs = hist[hist > 0] / max(float(hist.sum()), 1.0)
    return float(-(probs * np.log2(probs)).sum())


def _visual_metrics(blob: bytes, display_w_px: float, display_h_px: float) -> dict[str, Any]:
    image = Image.open(io.BytesIO(blob)).convert("RGB")
    rgb = np.asarray(image)
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    src_h, src_w = gray.shape

    # Normalize edge measurement so source resolution does not dominate it.
    scale_for_edges = min(1.0, 768.0 / max(src_w, src_h))
    edge_gray = cv2.resize(
        gray,
        (max(1, round(src_w * scale_for_edges)), max(1, round(src_h * scale_for_edges))),
        interpolation=cv2.INTER_AREA,
    )
    edges = cv2.Canny(edge_gray, 80, 180)
    edge_density = float(np.count_nonzero(edges) / edges.size)

    # Dark connected components approximate character strokes and small diagram
    # details. Boxes/lines are not fully separable without OCR, so retain raw
    # values for calibration rather than treating them as labels.
    _, binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    count, _, stats, _ = cv2.connectedComponentsWithStats(binary, connectivity=8)
    heights: list[int] = []
    max_component_area = max(12, int(src_w * src_h * 0.02))
    max_component_height = max(12, int(src_h * 0.12))
    for x, y, width, height, area in stats[1:count]:
        del x, y
        if not (2 <= height <= max_component_height):
            continue
        if not (1 <= width <= src_w * 0.5):
            continue
        if not (3 <= area <= max_component_area):
            continue
        aspect = width / max(height, 1)
        if 0.05 <= aspect <= 25:
            heights.append(int(height))

    fit_scale = min(display_w_px / max(src_w, 1), display_h_px / max(src_h, 1))
    component_heights_display = np.asarray(heights, dtype=float) * fit_scale
    median_detail = float(np.median(component_heights_display)) if heights else None
    p25_detail = float(np.percentile(component_heights_display, 25)) if heights else None

    return {
        "source_width_px": src_w,
        "source_height_px": src_h,
        "source_entropy": _round(_entropy(gray)),
        "edge_density": _round(edge_density, 4),
        "detail_component_count": len(heights),
        "median_detail_height_1080px": _round(median_detail, 2),
        "p25_detail_height_1080px": _round(p25_detail, 2),
        "source_to_display_scale": _round(fit_scale, 4),
    }


def _risk_flags(row: dict[str, Any]) -> str:
    flags: list[str] = []
    short_side = float(row["display_short_side_1080px"])
    area = float(row["display_area_1080px2"])
    edge_density = float(row["edge_density"])
    detail = row["median_detail_height_1080px"]
    if short_side < 160:
        flags.append("tiny_display_box")
    if area < 40000:
        flags.append("small_display_area")
    if detail != "" and float(detail) < 8:
        flags.append("small_internal_detail")
    if edge_density > 0.12:
        flags.append("high_visual_complexity")
    if edge_density > 0.08 and area < 80000:
        flags.append("complex_figure_in_small_box")
    return ",".join(flags)


def analyze(labels_path: Path, output_dir: Path) -> tuple[Path, Path]:
    with labels_path.open(newline="", encoding="utf-8") as fh:
        labels = [r for r in csv.DictReader(fh) if (r.get("primary_issue") or "").strip()]

    picture_rows: list[dict[str, Any]] = []
    poster_rows: list[dict[str, Any]] = []
    for label in labels:
        pptx_path = Path(label["iter1_png"].replace("iter_1.png", "iter_1.pptx"))
        prs = Presentation(pptx_path)
        slide_w_in = prs.slide_width / EMU_PER_INCH
        slide_h_in = prs.slide_height / EMU_PER_INCH
        px_per_in_x = TARGET_WIDTH / slide_w_in
        px_per_in_y = TARGET_HEIGHT / slide_h_in
        current: list[dict[str, Any]] = []

        for picture_index, shape in enumerate(
            (s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE), start=1
        ):
            display_w_in = shape.width / EMU_PER_INCH
            display_h_in = shape.height / EMU_PER_INCH
            display_w_px = display_w_in * px_per_in_x
            display_h_px = display_h_in * px_per_in_y
            row: dict[str, Any] = {
                "run_folder": label["run_folder"],
                "template": label["template"],
                "human_primary": label["primary_issue"],
                "human_secondary": label.get("secondary_issues", ""),
                "human_notes": label.get("free_notes", ""),
                "picture_index": picture_index,
                "display_width_in": _round(display_w_in),
                "display_height_in": _round(display_h_in),
                "display_width_1080px": _round(display_w_px, 1),
                "display_height_1080px": _round(display_h_px, 1),
                "display_short_side_1080px": _round(min(display_w_px, display_h_px), 1),
                "display_area_1080px2": _round(display_w_px * display_h_px, 1),
            }
            row.update(_visual_metrics(shape.image.blob, display_w_px, display_h_px))
            row["provisional_risk_flags"] = _risk_flags(row)
            current.append(row)
            picture_rows.append(row)

        numeric = lambda key: [float(r[key]) for r in current if r[key] != ""]
        poster_rows.append({
            "run_folder": label["run_folder"],
            "template": label["template"],
            "human_primary": label["primary_issue"],
            "human_secondary": label.get("secondary_issues", ""),
            "human_notes": label.get("free_notes", ""),
            "picture_count": len(current),
            "min_display_short_side_1080px": _round(min(numeric("display_short_side_1080px")), 1) if current else "",
            "min_display_area_1080px2": _round(min(numeric("display_area_1080px2")), 1) if current else "",
            "min_median_detail_height_1080px": _round(min(numeric("median_detail_height_1080px")), 2)
            if numeric("median_detail_height_1080px") else "",
            "max_edge_density": _round(max(numeric("edge_density")), 4) if current else "",
            "pictures_with_risk_flags": sum(bool(r["provisional_risk_flags"]) for r in current),
        })

    output_dir.mkdir(parents=True, exist_ok=True)
    pictures_path = output_dir / "picture_metrics.csv"
    posters_path = output_dir / "poster_summary.csv"
    for path, rows in ((pictures_path, picture_rows), (posters_path, poster_rows)):
        if not rows:
            continue
        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    return pictures_path, posters_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--labels", type=Path, default=DEFAULT_LABELS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    pictures, posters = analyze(args.labels, args.output_dir)
    print(f"picture metrics: {pictures}")
    print(f"poster summary: {posters}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
