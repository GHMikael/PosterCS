"""Deterministic poster measurements for T8-A.

This module deliberately emits continuous measurements and localized evidence.
It does not contain defect thresholds or produce defect labels.
"""

from __future__ import annotations

import io
import math
import re
from pathlib import Path
from typing import Any, Iterable

import cv2
import numpy as np
import pytesseract
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pytesseract import Output


FEATURE_VERSION = "t8a_feature_v1"
CANONICAL_WIDTH = 1920.0
CANONICAL_HEIGHT = 1080.0
GRID_COLS = 8
GRID_ROWS = 4
GRID_BLANK_OCCUPANCY = 0.15
BODY_TOP_RATIO = 0.12
BODY_BOTTOM_RATIO = 0.96


def _safe_div(a: float, b: float) -> float | None:
    return float(a / b) if b else None


def _percentile(values: Iterable[float], q: float) -> float | None:
    vals = [float(v) for v in values if v is not None and math.isfinite(float(v))]
    return float(np.percentile(vals, q)) if vals else None


def _cv(values: Iterable[float]) -> float | None:
    vals = np.asarray([float(v) for v in values if v is not None], dtype=float)
    if vals.size < 2 or float(vals.mean()) == 0:
        return None
    return float(vals.std() / vals.mean())


def _largest_rectangle(binary: np.ndarray) -> int:
    """Largest all-true axis-aligned rectangle in a boolean grid, in cells."""
    if binary.size == 0:
        return 0
    heights = np.zeros(binary.shape[1], dtype=int)
    best = 0
    for row in binary:
        heights = np.where(row, heights + 1, 0)
        stack: list[tuple[int, int]] = []
        for i, height in enumerate(np.append(heights, 0)):
            start = i
            while stack and stack[-1][1] > height:
                idx, h = stack.pop()
                best = max(best, h * (i - idx))
                start = idx
            if not stack or stack[-1][1] < height:
                stack.append((start, int(height)))
    return int(best)


def _bbox_iou(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    x1, y1 = max(ax, bx), max(ay, by)
    x2, y2 = min(ax + aw, bx + bw), min(ay + ah, by + bh)
    inter = max(0, x2 - x1) * max(0, y2 - y1)
    union = aw * ah + bw * bh - inter
    return inter / union if union else 0.0


def _cluster_residual(values: Iterable[float], tolerance: float) -> tuple[float | None, float | None, int]:
    vals = sorted(float(v) for v in values)
    if len(vals) < 2:
        return None, None, 0
    clusters: list[list[float]] = []
    for value in vals:
        if not clusters or value - float(np.median(clusters[-1])) > tolerance:
            clusters.append([value])
        else:
            clusters[-1].append(value)
    residuals: list[float] = []
    used = 0
    for cluster in clusters:
        if len(cluster) < 2:
            continue
        center = float(np.median(cluster))
        residuals.extend(abs(v - center) for v in cluster)
        used += len(cluster)
    return _percentile(residuals, 50), max(residuals) if residuals else None, used


def _foreground_mask(rgb: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    hsv = cv2.cvtColor(rgb, cv2.COLOR_RGB2HSV)
    edges = cv2.Canny(gray, 60, 160)
    # Visual-ink mask: dark pixels, saturated pixels, and crisp edges. This is
    # robust to posters that use several different near-white panel backgrounds.
    mask = ((gray < 225) | (hsv[:, :, 1] > 30) | (edges > 0)).astype(np.uint8)
    kernel = np.ones((3, 3), np.uint8)
    return cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel)


def _grid_occupancy(mask: np.ndarray) -> np.ndarray:
    h, w = mask.shape
    result = np.zeros((GRID_ROWS, GRID_COLS), dtype=float)
    for row in range(GRID_ROWS):
        for col in range(GRID_COLS):
            y1, y2 = round(row * h / GRID_ROWS), round((row + 1) * h / GRID_ROWS)
            x1, x2 = round(col * w / GRID_COLS), round((col + 1) * w / GRID_COLS)
            result[row, col] = float(mask[y1:y2, x1:x2].mean())
    return result


def _ocr(rgb: np.ndarray) -> list[dict[str, Any]]:
    h, w = rgb.shape[:2]
    scale = min(1.0, 1800.0 / w)
    work = cv2.resize(rgb, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA) if scale < 1 else rgb
    data = pytesseract.image_to_data(work, config="--psm 11", output_type=Output.DICT)
    words: list[dict[str, Any]] = []
    for i, raw in enumerate(data.get("text", [])):
        text = str(raw).strip()
        try:
            conf = float(data["conf"][i])
        except (ValueError, TypeError):
            conf = -1.0
        if not text or conf < 0:
            continue
        x = int(round(int(data["left"][i]) / scale))
        y = int(round(int(data["top"][i]) / scale))
        bw = int(round(int(data["width"][i]) / scale))
        bh = int(round(int(data["height"][i]) / scale))
        words.append({"text": text, "confidence": conf, "x": x, "y": y, "w": bw, "h": bh,
                      "block": int(data["block_num"][i]), "paragraph": int(data["par_num"][i]),
                      "line": int(data["line_num"][i])})
    return words


def _ocr_lines(words: list[dict[str, Any]], image_w: int, image_h: int) -> list[dict[str, Any]]:
    grouped: dict[tuple[int, int, int], list[dict[str, Any]]] = {}
    for word in words:
        grouped.setdefault((word["block"], word["paragraph"], word["line"]), []).append(word)
    lines: list[dict[str, Any]] = []
    for line_words in grouped.values():
        line_words.sort(key=lambda x: x["x"])
        x1 = min(x["x"] for x in line_words)
        y1 = min(x["y"] for x in line_words)
        x2 = max(x["x"] + x["w"] for x in line_words)
        y2 = max(x["y"] + x["h"] for x in line_words)
        lines.append({
            "region_type": "ocr_line", "x_norm": x1 / image_w, "y_norm": y1 / image_h,
            "w_norm": (x2 - x1) / image_w, "h_norm": (y2 - y1) / image_h,
            "text": " ".join(x["text"] for x in line_words),
            "confidence": float(np.mean([x["confidence"] for x in line_words])),
            "height_px_1080": (y2 - y1) * CANONICAL_HEIGHT / image_h,
        })
    return sorted(lines, key=lambda x: (x["y_norm"], x["x_norm"]))


def _panel_candidates(rgb: np.ndarray) -> list[tuple[int, int, int, int]]:
    h, w = rgb.shape[:2]
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    edges = cv2.Canny(gray, 40, 120)
    edges = cv2.morphologyEx(edges, cv2.MORPH_CLOSE, np.ones((5, 5), np.uint8))
    contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    boxes: list[tuple[int, int, int, int]] = []
    for contour in contours:
        x, y, bw, bh = cv2.boundingRect(contour)
        area_ratio = bw * bh / (w * h)
        if 0.012 <= area_ratio <= 0.48 and bw >= 0.08 * w and bh >= 0.08 * h:
            boxes.append((x, y, bw, bh))
    boxes.sort(key=lambda b: b[2] * b[3], reverse=True)
    kept: list[tuple[int, int, int, int]] = []
    for box in boxes:
        if all(_bbox_iou(box, prior) < 0.82 for prior in kept):
            kept.append(box)
    return sorted(kept[:24], key=lambda b: (b[1], b[0]))


def _component_heights(rgb: np.ndarray) -> list[float]:
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    edge = cv2.Canny(gray, 60, 160)
    count, _, stats, _ = cv2.connectedComponentsWithStats(edge, 8)
    image_area = rgb.shape[0] * rgb.shape[1]
    return [float(stats[i, cv2.CC_STAT_HEIGHT]) for i in range(1, count)
            if 4 <= stats[i, cv2.CC_STAT_AREA] <= image_area * 0.01]


def extract_raster(image_path: str | Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    path = Path(image_path)
    rgb = np.asarray(Image.open(path).convert("RGB"))
    h, w = rgb.shape[:2]
    mask = _foreground_mask(rgb)
    grid = _grid_occupancy(mask)
    words = _ocr(rgb)
    lines = _ocr_lines(words, w, h)
    body_line_heights = [x["height_px_1080"] for x in lines
                         if BODY_TOP_RATIO <= x["y_norm"] <= BODY_BOTTOM_RATIO]
    body_words = [x for x in words if BODY_TOP_RATIO * h <= x["y"] <= BODY_BOTTOM_RATIO * h]
    body_heights = [x["h"] * CANONICAL_HEIGHT / h for x in body_words]
    boxes = _panel_candidates(rgb)
    large_panel_boxes = [b for b in boxes if b[2] * b[3] / (w * h) >= 0.07 and b[1] / h >= BODY_TOP_RATIO
                         and b[3] / h >= 0.30]
    comp_h = [x * CANONICAL_HEIGHT / h for x in _component_heights(rgb)]
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    edge = cv2.Canny(gray, 60, 160)
    left = float(mask[:, : w // 2].mean())
    right = float(mask[:, w // 2 :].mean())
    top = float(mask[: h // 2].mean())
    bottom = float(mask[h // 2 :].mean())
    edge_touch = sum(1 for x in words if x["x"] <= 0.005 * w or x["x"] + x["w"] >= 0.995 * w
                     or x["y"] <= 0.005 * h or x["y"] + x["h"] >= 0.995 * h)
    ellipsis = sum(x["text"].count("...") + x["text"].count("…") for x in words)
    chars = sum(len(re.sub(r"\s+", "", x["text"])) for x in words)
    text_area = sum(x["w"] * x["h"] for x in words)
    med_align, max_align, align_used = _cluster_residual(
        [v for b in boxes for v in (b[0] / w, (b[0] + b[2]) / w)], tolerance=0.015
    )
    large_x_med, large_x_max, large_x_used = _cluster_residual(
        [v for b in large_panel_boxes for v in (b[0] / w, (b[0] + b[2]) / w)], tolerance=0.05
    )
    large_y_med, large_y_max, large_y_used = _cluster_residual(
        [v for b in large_panel_boxes for v in (b[1] / h, (b[1] + b[3]) / h)], tolerance=0.05
    )
    features: dict[str, Any] = {
        "raster_available": 1, "image_width_px": w, "image_height_px": h,
        "image_aspect_ratio": w / h, "foreground_ratio": float(mask.mean()),
        "foreground_left_ratio": left, "foreground_right_ratio": right,
        "foreground_left_right_abs_diff": abs(left - right),
        "foreground_top_ratio": top, "foreground_bottom_ratio": bottom,
        "foreground_top_bottom_abs_diff": abs(top - bottom),
        "grid_occupancy_mean": float(grid.mean()), "grid_occupancy_std": float(grid.std()),
        "grid_occupancy_min": float(grid.min()), "grid_occupancy_max": float(grid.max()),
        "grid_occupancy_p10": _percentile(grid.ravel(), 10),
        "grid_occupancy_p25": _percentile(grid.ravel(), 25),
        "largest_blank_grid_rect_ratio": _largest_rectangle(grid < GRID_BLANK_OCCUPANCY) / grid.size,
        "largest_very_blank_grid_rect_ratio": _largest_rectangle(grid < 0.10) / grid.size,
        "largest_low_density_grid_rect_ratio": _largest_rectangle(grid < 0.20) / grid.size,
        "edge_density": float((edge > 0).mean()), "ocr_word_count": len(words),
        "ocr_line_count": len(lines), "ocr_character_count": chars,
        "ocr_mean_confidence": float(np.mean([x["confidence"] for x in words])) if words else None,
        "ocr_text_bbox_area_ratio": text_area / (w * h),
        "ocr_characters_per_megapixel": chars / (w * h / 1_000_000),
        "ocr_characters_per_text_bbox_area": _safe_div(chars, text_area),
        "body_word_height_p10_px1080": _percentile(body_heights, 10),
        "body_word_height_p25_px1080": _percentile(body_heights, 25),
        "body_word_height_median_px1080": _percentile(body_heights, 50),
        "body_line_height_p10_px1080": _percentile(body_line_heights, 10),
        "body_line_height_p25_px1080": _percentile(body_line_heights, 25),
        "body_line_height_median_px1080": _percentile(body_line_heights, 50),
        "component_height_p10_px1080": _percentile(comp_h, 10),
        "component_height_p25_px1080": _percentile(comp_h, 25),
        "component_height_median_px1080": _percentile(comp_h, 50),
        "ellipsis_token_count": ellipsis, "edge_touching_ocr_word_count": edge_touch,
        "panel_candidate_count": len(boxes),
        "panel_width_cv": _cv([b[2] / w for b in boxes]), "panel_height_cv": _cv([b[3] / h for b in boxes]),
        "panel_edge_alignment_median_residual_norm": med_align,
        "panel_edge_alignment_max_residual_norm": max_align,
        "panel_edges_in_clusters": align_used,
        "large_panel_candidate_count": len(large_panel_boxes),
        "large_panel_grid_alignment_median_residual_norm": max(
            x for x in (large_x_med, large_y_med) if x is not None
        ) if any(x is not None for x in (large_x_med, large_y_med)) else None,
        "large_panel_grid_alignment_max_residual_norm": max(
            x for x in (large_x_max, large_y_max) if x is not None
        ) if any(x is not None for x in (large_x_max, large_y_max)) else None,
        "large_panel_edges_in_clusters": large_x_used + large_y_used,
    }
    regions = list(lines)
    regions.extend({"region_type": "panel_candidate", "x_norm": x / w, "y_norm": y / h,
                    "w_norm": bw / w, "h_norm": bh / h, "area_ratio": bw * bh / (w * h)}
                   for x, y, bw, bh in boxes)
    return features, regions


def _shape_bbox(shape: Any, slide_w: int, slide_h: int) -> tuple[float, float, float, float]:
    return shape.left / slide_w, shape.top / slide_h, shape.width / slide_w, shape.height / slide_h


def _font_sizes(shape: Any) -> tuple[list[float], int, int, int]:
    sizes: list[float] = []
    runs = unresolved = paragraphs = bullets = 0
    if not getattr(shape, "has_text_frame", False):
        return sizes, runs, unresolved, paragraphs
    for paragraph in shape.text_frame.paragraphs:
        paragraphs += 1
        if paragraph.level > 0:
            bullets += 1
        paragraph_xml_sizes: list[float] = []
        for element in paragraph._p.xpath(".//*[@sz]"):
            try:
                paragraph_xml_sizes.append(int(element.get("sz")) / 100.0)
            except (TypeError, ValueError):
                pass
        inherited_size = float(np.median(paragraph_xml_sizes)) if paragraph_xml_sizes else None
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            runs += 1
            if run.font.size is None:
                if inherited_size is None:
                    unresolved += 1
                else:
                    sizes.append(inherited_size)
            else:
                sizes.append(float(run.font.size.pt))
    return sizes, runs, unresolved, bullets


def _embedded_image_metrics(blob: bytes, display_w_px: float, display_h_px: float) -> dict[str, Any]:
    try:
        rgb = np.asarray(Image.open(io.BytesIO(blob)).convert("RGB"))
    except Exception:
        return {"decode_ok": 0}
    h, w = rgb.shape[:2]
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    edge = cv2.Canny(gray, 60, 160)
    comp = _component_heights(rgb)
    display_scale = min(display_w_px / w, display_h_px / h) if w and h else None
    return {
        "decode_ok": 1, "source_width_px": w, "source_height_px": h,
        "display_short_side_px1080": min(display_w_px, display_h_px),
        "display_area_ratio": display_w_px * display_h_px / (CANONICAL_WIDTH * CANONICAL_HEIGHT),
        "source_to_display_scale": display_scale, "embedded_edge_density": float((edge > 0).mean()),
        "detail_height_p25_display_px": (_percentile(comp, 25) * display_scale) if comp and display_scale else None,
        "detail_height_median_display_px": (_percentile(comp, 50) * display_scale) if comp and display_scale else None,
    }


def _intersection_ratio(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    inter = max(0.0, min(ax + aw, bx + bw) - max(ax, bx)) * max(0.0, min(ay + ah, by + bh) - max(ay, by))
    return inter / min(aw * ah, bw * bh) if aw * ah > 0 and bw * bh > 0 else 0.0


def _box_occupancy(boxes: list[tuple[float, float, float, float]], rows: int = 40, cols: int = 80) -> np.ndarray:
    """Rasterize normalized source boxes into a deterministic occupancy mask."""
    mask = np.zeros((rows, cols), dtype=np.uint8)
    for x, y, w, h in boxes:
        x1 = max(0, min(cols, int(math.floor(x * cols))))
        x2 = max(0, min(cols, int(math.ceil((x + w) * cols))))
        y1 = max(0, min(rows, int(math.floor(y * rows))))
        y2 = max(0, min(rows, int(math.ceil((y + h) * rows))))
        if x2 > x1 and y2 > y1:
            mask[y1:y2, x1:x2] = 1
    return mask


def _coarse_occupancy(mask: np.ndarray) -> np.ndarray:
    h, w = mask.shape
    grid = np.zeros((GRID_ROWS, GRID_COLS), dtype=float)
    for row in range(GRID_ROWS):
        for col in range(GRID_COLS):
            y1, y2 = round(row * h / GRID_ROWS), round((row + 1) * h / GRID_ROWS)
            x1, x2 = round(col * w / GRID_COLS), round((col + 1) * w / GRID_COLS)
            grid[row, col] = float(mask[y1:y2, x1:x2].mean())
    return grid


def extract_source(pptx_path: str | Path | None) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    if not pptx_path or not Path(pptx_path).exists():
        return {"source_available": 0}, []
    prs = Presentation(str(pptx_path))
    slide = prs.slides[0]
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    features: dict[str, Any] = {"source_available": 1, "source_slide_width_emu": sw,
                                "source_slide_height_emu": sh, "source_shape_count": len(slide.shapes)}
    regions: list[dict[str, Any]] = []
    all_sizes: list[float] = []
    body_sizes: list[float] = []
    text_densities: list[float] = []
    picture_metrics: list[dict[str, Any]] = []
    content_boxes: list[tuple[float, float, float, float]] = []
    body_content_boxes: list[tuple[float, float, float, float]] = []
    panel_boxes: list[tuple[float, float, float, float]] = []
    named_panel_boxes: list[tuple[float, float, float, float]] = []
    text_count = picture_count = chars = paragraphs = bullets = unresolved = total_runs = out_of_bounds = 0
    ellipsis = 0
    for idx, shape in enumerate(slide.shapes):
        x, y, w, h = _shape_bbox(shape, sw, sh)
        if x < 0 or y < 0 or x + w > 1 or y + h > 1:
            out_of_bounds += 1
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture_count += 1
            content_boxes.append((x, y, w, h))
            if y >= BODY_TOP_RATIO and y + h <= BODY_BOTTOM_RATIO:
                body_content_boxes.append((x, y, w, h))
            pm = _embedded_image_metrics(shape.image.blob, w * CANONICAL_WIDTH, h * CANONICAL_HEIGHT)
            picture_metrics.append(pm)
            regions.append({"region_type": "source_picture", "shape_index": idx, "x_norm": x, "y_norm": y,
                            "w_norm": w, "h_norm": h, **pm})
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            text_count += 1
            text = shape.text.strip()
            clean_chars = len(re.sub(r"\s+", "", text))
            chars += clean_chars
            ellipsis += text.count("...") + text.count("…")
            sizes, runs, missing, bullet_count = _font_sizes(shape)
            all_sizes.extend(sizes)
            if y >= BODY_TOP_RATIO and y + h <= BODY_BOTTOM_RATIO:
                body_sizes.extend(sizes)
            total_runs += runs
            unresolved += missing
            paragraphs += len(shape.text_frame.paragraphs)
            bullets += bullet_count
            density = _safe_div(clean_chars, w * h)
            if density is not None:
                text_densities.append(density)
            content_boxes.append((x, y, w, h))
            if y >= BODY_TOP_RATIO and y + h <= BODY_BOTTOM_RATIO:
                body_content_boxes.append((x, y, w, h))
            regions.append({"region_type": "source_text", "shape_index": idx, "x_norm": x, "y_norm": y,
                            "w_norm": w, "h_norm": h, "character_count": clean_chars,
                            "paragraph_count": len(shape.text_frame.paragraphs), "font_min_pt": min(sizes) if sizes else None,
                            "font_median_pt": _percentile(sizes, 50), "characters_per_norm_area": density,
                            "text_preview": text[:160]})
        shape_name = str(getattr(shape, "name", ""))
        if re.match(r"^panel_\d+_", shape_name):
            named_panel_boxes.append((x, y, w, h))
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 0.01 <= w * h <= 0.5 and w >= 0.08 and h >= 0.08:
            panel_boxes.append((x, y, w, h))
            regions.append({"region_type": "source_panel_candidate", "shape_index": idx,
                            "shape_name": shape_name, "x_norm": x, "y_norm": y, "w_norm": w, "h_norm": h})
    intersections = [_intersection_ratio(a, b) for i, a in enumerate(content_boxes) for b in content_boxes[i + 1 :]]
    content_mask = _box_occupancy(body_content_boxes)
    content_grid = _coarse_occupancy(content_mask)
    source_left = float(content_mask[:, : content_mask.shape[1] // 2].mean())
    source_right = float(content_mask[:, content_mask.shape[1] // 2 :].mean())
    source_top = float(content_mask[: content_mask.shape[0] // 2].mean())
    source_bottom = float(content_mask[content_mask.shape[0] // 2 :].mean())
    alignment_values = [v for b in panel_boxes for v in (b[0], b[0] + b[2])]
    med_align, max_align, align_used = _cluster_residual(alignment_values, 0.015)
    named_x_med, named_x_max, named_x_used = _cluster_residual(
        [v for b in named_panel_boxes for v in (b[0], b[0] + b[2])], 0.05
    )
    named_y_med, named_y_max, named_y_used = _cluster_residual(
        [v for b in named_panel_boxes for v in (b[1], b[1] + b[3])], 0.05
    )
    picture_short = [m.get("display_short_side_px1080") for m in picture_metrics if m.get("decode_ok")]
    picture_scale = [m.get("source_to_display_scale") for m in picture_metrics if m.get("decode_ok")]
    picture_detail = [m.get("detail_height_p25_display_px") for m in picture_metrics if m.get("detail_height_p25_display_px") is not None]
    features.update({
        "source_text_shape_count": text_count, "source_picture_count": picture_count,
        "source_character_count": chars, "source_paragraph_count": paragraphs, "source_bullet_level_count": bullets,
        "source_ellipsis_token_count": ellipsis, "source_out_of_bounds_shape_count": out_of_bounds,
        "source_font_run_count": total_runs, "source_unresolved_font_run_count": unresolved,
        "source_font_resolved_ratio": _safe_div(total_runs - unresolved, total_runs),
        "source_font_min_pt": min(all_sizes) if all_sizes else None,
        "source_font_p10_pt": _percentile(all_sizes, 10), "source_font_median_pt": _percentile(all_sizes, 50),
        "source_body_font_min_pt": min(body_sizes) if body_sizes else None,
        "source_body_font_p10_pt": _percentile(body_sizes, 10),
        "source_body_font_median_pt": _percentile(body_sizes, 50),
        "source_text_density_median": _percentile(text_densities, 50),
        "source_text_density_max": max(text_densities) if text_densities else None,
        "source_picture_short_side_min_px1080": min(picture_short) if picture_short else None,
        "source_picture_short_side_median_px1080": _percentile(picture_short, 50),
        "source_picture_scale_min": min(picture_scale) if picture_scale else None,
        "source_picture_detail_p25_min_display_px": min(picture_detail) if picture_detail else None,
        "source_panel_candidate_count": len(panel_boxes),
        "source_panel_width_cv": _cv([b[2] for b in panel_boxes]),
        "source_panel_height_cv": _cv([b[3] for b in panel_boxes]),
        "source_panel_edge_alignment_median_residual_norm": med_align,
        "source_panel_edge_alignment_max_residual_norm": max_align,
        "source_panel_edges_in_clusters": align_used,
        "source_named_panel_count": len(named_panel_boxes),
        "source_named_panel_grid_alignment_median_residual_norm": max(
            x for x in (named_x_med, named_y_med) if x is not None
        ) if any(x is not None for x in (named_x_med, named_y_med)) else None,
        "source_named_panel_grid_alignment_max_residual_norm": max(
            x for x in (named_x_max, named_y_max) if x is not None
        ) if any(x is not None for x in (named_x_max, named_y_max)) else None,
        "source_named_panel_edges_in_clusters": named_x_used + named_y_used,
        "source_content_overlap_pair_count": sum(v > 0 for v in intersections),
        "source_content_overlap_max_min_area_ratio": max(intersections) if intersections else 0.0,
        "source_content_shape_occupancy_ratio": float(content_mask.mean()),
        "source_content_left_occupancy": source_left,
        "source_content_right_occupancy": source_right,
        "source_content_left_right_abs_diff": abs(source_left - source_right),
        "source_content_top_occupancy": source_top,
        "source_content_bottom_occupancy": source_bottom,
        "source_content_top_bottom_abs_diff": abs(source_top - source_bottom),
        "source_content_grid_std": float(content_grid.std()),
        "source_largest_empty_grid_rect_ratio": _largest_rectangle(content_grid == 0) / content_grid.size,
    })
    return features, regions


__all__ = [
    "FEATURE_VERSION", "GRID_BLANK_OCCUPANCY", "BODY_TOP_RATIO", "BODY_BOTTOM_RATIO",
    "extract_raster", "extract_source", "_largest_rectangle", "_cluster_residual", "_box_occupancy",
]
