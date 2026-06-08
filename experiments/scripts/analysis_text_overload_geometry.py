#!/usr/bin/env python3
"""Track B: does an AREA-AWARE proxy (chars/cm^2) + real rendered font size ground
`text_overload` where raw char-count failed?

analysis_text_overload_validity.py found AUC(text_overload>visual_hierarchy_weak) = 0.637 on
mean_chars_per_panel — barely above coin flip. But mean_chars_per_panel ignores how large the
text box actually is: the SAME char count in a 4 cm^2 box vs a 20 cm^2 box is wildly different
visual density. The PPTX stores the true per-shape geometry (EMU->cm) and the explicitly set
font sizes (pt). This script mines iter_1.pptx for each poster and re-tests separation with:

  - chars_per_cm2_median : median over content text boxes of (chars / box_area)  [headline]
  - chars_per_cm2_max    : densest single content box
  - chars_per_cm2_agg    : sum(chars) / sum(area) over content boxes
  - text_area_fraction   : sum(text-box area) / slide area
  - body_font_pt_min/med : explicitly-set run font sizes (the small body text)

A "content box" = a text box with >= MIN_CHARS chars (drops 1-char bullet dots / decorations).
Zero API. Reads labels/*.json (for the VLM primary + run_dir) and each run's pptx/iter_1.pptx.
"""
from __future__ import annotations

import argparse
import glob
import json
import statistics
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

OK_SOURCES = {"vlm", "vlm_partial", "human"}
BLIP_RUN_ID = "44a0517194a2"
MIN_CHARS = 10  # a text box must hold this many chars to count as "content" (drops bullet dots)

GEOM_KEYS = [
    "chars_per_cm2_median",
    "chars_per_cm2_max",
    "chars_per_cm2_agg",
    "text_area_fraction",
    "body_font_pt_min",
    "body_font_pt_median",
    "n_content_boxes",
]
# Direction note: for the first four, HIGHER = more overloaded. For font sizes, LOWER = more
# overloaded (smaller text crammed in). We report AUC(TO>VHW) for all and interpret accordingly.
HEADLINE_KEYS = ["chars_per_cm2_median", "chars_per_cm2_agg", "text_area_fraction", "body_font_pt_min"]


def chosen_label(rec: dict) -> dict | None:
    labels = rec.get("labels") or []
    if not labels:
        return None
    for lab in labels:
        if lab.get("source") == "human":
            return lab
    return labels[0]


def extract_geom(pptx_path: Path) -> dict | None:
    """Per-slide geometry/font proxies from the first slide of a poster pptx."""
    try:
        prs = Presentation(str(pptx_path))
    except Exception:
        return None
    if not prs.slides:
        return None
    sl = prs.slides[0]
    try:
        slide_area = prs.slide_width.cm * prs.slide_height.cm
    except Exception:
        slide_area = None

    box_densities: list[float] = []
    content_chars = 0
    content_area = 0.0
    n_content = 0
    all_text_area = 0.0
    fonts: list[float] = []

    for sh in sl.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        tf = sh.text_frame
        txt = "".join(p.text for p in tf.paragraphs)
        n = len(txt.strip())
        try:
            area = sh.width.cm * sh.height.cm
        except Exception:
            area = None
        if area is None or area <= 0:
            continue
        all_text_area += area
        # Collect explicitly-set font sizes (run-level, else paragraph-level).
        for p in tf.paragraphs:
            psz = p.font.size
            runs = p.runs or []
            if not runs and psz is not None:
                fonts.append(psz.pt)
            for run in runs:
                sz = run.font.size or psz
                if sz is not None:
                    fonts.append(sz.pt)
        if n >= MIN_CHARS:
            content_chars += n
            content_area += area
            n_content += 1
            box_densities.append(n / area)

    def med(xs):
        return round(statistics.median(xs), 3) if xs else None

    return {
        "slide_area_cm2": round(slide_area, 1) if slide_area else None,
        "chars_per_cm2_median": med(box_densities),
        "chars_per_cm2_max": round(max(box_densities), 3) if box_densities else None,
        "chars_per_cm2_agg": round(content_chars / content_area, 3) if content_area else None,
        "text_area_fraction": round(all_text_area / slide_area, 3) if slide_area else None,
        "body_font_pt_min": round(min(fonts), 1) if fonts else None,
        "body_font_pt_median": med(fonts),
        "n_content_boxes": n_content,
        "n_font_samples": len(fonts),
    }


def describe(values):
    vals = [v for v in values if v is not None]
    if not vals:
        return None
    n = len(vals)
    return {
        "n": n,
        "mean": round(statistics.mean(vals), 2),
        "median": round(statistics.median(vals), 2),
        "std": round(statistics.pstdev(vals), 2) if n > 1 else 0.0,
        "min": round(min(vals), 2),
        "max": round(max(vals), 2),
    }


def auc(pos, neg):
    pos = [v for v in pos if v is not None]
    neg = [v for v in neg if v is not None]
    if not pos or not neg:
        return None
    wins = 0.0
    for a in pos:
        for b in neg:
            wins += 1.0 if a > b else (0.5 if a == b else 0.0)
    return round(wins / (len(pos) * len(neg)), 3)


def pct_rank(value, pool):
    pool = [v for v in pool if v is not None]
    if value is None or not pool:
        return None
    below = sum(1 for v in pool if v < value)
    return round(100.0 * below / len(pool), 1)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--labels-dir", default="experiments/results/audit/labels")
    ap.add_argument("--out", default="experiments/results/audit/analysis_text_overload_geometry.json")
    ap.add_argument("--pptx-name", default="iter_1.pptx")
    args = ap.parse_args()

    items = []
    skipped = []
    for fp in sorted(glob.glob(str(Path(args.labels_dir) / "*.json"))):
        rec = json.loads(Path(fp).read_text())
        lab = chosen_label(rec)
        if lab is None or lab.get("source") not in OK_SOURCES or not lab.get("primary_issue"):
            continue
        run_dir = rec.get("run_dir") or ""
        pptx_path = Path(run_dir) / "pptx" / args.pptx_name
        if not pptx_path.exists():
            skipped.append((rec.get("run_id"), "no pptx"))
            continue
        geom = extract_geom(pptx_path)
        if geom is None or geom.get("chars_per_cm2_median") is None:
            skipped.append((rec.get("run_id"), "no geom"))
            continue
        items.append(
            {
                "run_id": rec.get("run_id"),
                "template": rec.get("template"),
                "primary": lab.get("primary_issue"),
                "secondary": list(lab.get("secondary_issues") or []),
                "mean_chars_per_panel": (rec.get("proxies") or {}).get("mean_chars_per_panel"),
                "geom": geom,
            }
        )

    by_primary = defaultdict(list)
    for it in items:
        by_primary[it["primary"]].append(it)
    primary_counts = {k: len(v) for k, v in sorted(by_primary.items(), key=lambda kv: -len(kv[1]))}

    def vals(group, key):
        return [g["geom"].get(key) for g in group if g["geom"].get(key) is not None]

    group_stats = {p: {k: describe(vals(g, k)) for k in GEOM_KEYS} for p, g in by_primary.items()}

    TO, VHW = "text_overload", "visual_hierarchy_weak"
    to_g, vhw_g = by_primary.get(TO, []), by_primary.get(VHW, [])
    separation = {}
    for k in GEOM_KEYS:
        separation[k] = {
            "auc_TO_vs_VHW": auc(vals(to_g, k), vals(vhw_g, k)),
            "TO_median": (describe(vals(to_g, k)) or {}).get("median"),
            "VHW_median": (describe(vals(vhw_g, k)) or {}).get("median"),
        }

    blip = next((it for it in items if it["run_id"] == BLIP_RUN_ID), None)
    blip_loc = None
    if blip:
        blip_loc = {"run_id": BLIP_RUN_ID, "primary": blip["primary"], "geom": blip["geom"]}
        for k in GEOM_KEYS:
            blip_loc[f"pct_{k}"] = pct_rank(blip["geom"].get(k), vals(items, k))

    head = separation["chars_per_cm2_median"]["auc_TO_vs_VHW"]
    if head is None:
        verdict = "INDETERMINATE — empty group."
    elif head >= 0.70:
        verdict = ("GROUNDED-BY-AREA — chars/cm^2 separates text_overload from visual_hierarchy_weak "
                   "where raw char-count did not. The construct is real once box area is accounted for.")
    elif head <= 0.60:
        verdict = ("STILL PRIOR-LIKE — even area-aware density does not separate the two. text_overload "
                   "is applied largely independent of measurable density.")
    else:
        verdict = "AMBIGUOUS — area-aware density helps only weakly; small n_VHW."

    result = {
        "n": len(items),
        "n_skipped": len(skipped),
        "skipped": skipped,
        "min_chars_per_content_box": MIN_CHARS,
        "primary_counts": primary_counts,
        "group_geom_stats": group_stats,
        "separation_TO_vs_VHW": separation,
        "baseline_auc_mean_chars_per_panel": 0.637,
        "headline_auc_chars_per_cm2_median": head,
        "blip_locator": blip_loc,
        "verdict": verdict,
    }
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False))

    # ---- console ----
    print(f"[geometry] n={len(items)} skipped={len(skipped)}  primary={primary_counts}")
    print(f"[geometry] wrote {out}")
    print()
    print("== geometry by PRIMARY group (median [min..max], n) ==")
    for p in sorted(by_primary, key=lambda k: -len(by_primary[k])):
        parts = [f"{p:<26} n={len(by_primary[p]):<2}"]
        for k in ["chars_per_cm2_median", "text_area_fraction", "body_font_pt_min"]:
            d = group_stats[p][k]
            if d:
                parts.append(f"{k}={d['median']}[{d['min']}..{d['max']}]")
        print("  " + "  ".join(parts))
    print()
    print("== AUC(text_overload > visual_hierarchy_weak) — higher chars/cm2 & area = more overload ==")
    print(f"  (baseline mean_chars_per_panel AUC = 0.637)")
    for k in HEADLINE_KEYS:
        s = separation[k]
        note = "  <- font: LOWER=overload, so <0.5 supports TO" if "font" in k else ""
        print(f"  {k:<24} AUC={s['auc_TO_vs_VHW']}   TO_med={s['TO_median']}  VHW_med={s['VHW_median']}{note}")
    print()
    if blip_loc:
        g = blip_loc["geom"]
        print("== BLIP (the poster the user judged '刚好') ==")
        print(f"  primary={blip_loc['primary']}")
        print(f"  chars/cm2_median={g['chars_per_cm2_median']} (pct {blip_loc['pct_chars_per_cm2_median']}%)  "
              f"text_area_frac={g['text_area_fraction']} (pct {blip_loc['pct_text_area_fraction']}%)")
        print(f"  body_font_pt_min={g['body_font_pt_min']} (pct {blip_loc['pct_body_font_pt_min']}%)  "
              f"n_content_boxes={g['n_content_boxes']}")
    print()
    print("VERDICT:", verdict)


if __name__ == "__main__":
    main()
