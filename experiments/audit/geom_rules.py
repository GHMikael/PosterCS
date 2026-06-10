"""Geometric objective rules for the audit (negative controls / hallucination probes).

These are deterministic, rule-based detectors that run on the real generated
artifacts (the ``pptx/iter_1.pptx``), NOT VLM judgements. On PosterCS's own
clean posters they should fire near-zero (so a VLM that reports the matching
issue is making a false positive); on messy external SOTA output (e.g.
Paper2Poster) the same rules genuinely fire — demonstrating that geometric
detectors transfer while the holistic VLM critic is the unreliable link.
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Dict


# ---------------------------------------------------------------------------
# overlap rule — two element bounding boxes physically intersect.
# ---------------------------------------------------------------------------


def _overlap_ratio(a, b) -> float:
    """Intersection area over the smaller box's area, for boxes (l, t, w, h).

    Returns 0.0 when the boxes do not intersect or either has non-positive area.
    """
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    if aw <= 0 or ah <= 0 or bw <= 0 or bh <= 0:
        return 0.0
    ix = max(0.0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0.0, min(ay + ah, by + bh) - max(ay, by))
    inter = ix * iy
    if inter <= 0:
        return 0.0
    return inter / min(aw * ah, bw * bh)


def _is_collision(a, b, thresh: float = 0.15, containment_max: float = 0.9) -> bool:
    """True iff a and b *partially* overlap (a genuine collision).

    PPTX posters are deeply nested (page bg ⊃ panel cards ⊃ sub-cards / text /
    pictures). Normal nesting gives ``inter/min_area ≈ 1.0`` and must NOT count
    as a defect. A real overlap defect is a *partial* intersection where neither
    box contains the other: ``thresh ≤ inter/min_area < containment_max``.
    """
    r = _overlap_ratio(a, b)
    return thresh <= r < containment_max


def overlap_violations(pptx_path, thresh: float = 0.15) -> List[Dict]:
    """Pairs of shapes on slide 1 that *partially* overlap (genuine collisions).

    Uses python-pptx shape geometry (EMU). Containment (normal nesting) is
    excluded by :func:`_is_collision`. Returns ``{"i", "j", "ratio"}`` records
    (empty = clean). On a well-formed poster this should be ~empty; on messy
    output (overlapping/clipped elements) it fires.
    """
    from pptx import Presentation

    prs = Presentation(str(Path(pptx_path)))
    slide = prs.slides[0]
    boxes = []
    for s in slide.shapes:
        if s.left is None or s.top is None or not s.width or not s.height:
            continue
        boxes.append((int(s.left), int(s.top), int(s.width), int(s.height)))
    out: List[Dict] = []
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            if _is_collision(boxes[i], boxes[j], thresh=thresh):
                out.append({"i": i, "j": j, "ratio": round(_overlap_ratio(boxes[i], boxes[j]), 3)})
    return out


# ---------------------------------------------------------------------------
# contrast rule — theme-level WCAG check (deterministic negative control).
# Posters use a fixed palette per ``color_theme``; the rendered text/background
# colors are a property of the theme, so contrast is decided by the theme, not
# by per-pixel analysis. If every shipped theme is WCAG-compliant, then any VLM
# ``low_contrast`` report is a false positive (hallucination probe).
# ---------------------------------------------------------------------------

AA_NORMAL = 4.5   # WCAG 2.1 AA for normal text
AA_LARGE = 3.0    # WCAG 2.1 AA for large text (headers)


def _rel_luminance(rgb) -> float:
    """WCAG relative luminance for an (r, g, b) triple in 0..255."""
    def lin(v: float) -> float:
        v = v / 255.0
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = rgb[0], rgb[1], rgb[2]
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def wcag_ratio(rgb1, rgb2) -> float:
    """WCAG contrast ratio between two colors (1.0 .. 21.0)."""
    l1, l2 = _rel_luminance(rgb1), _rel_luminance(rgb2)
    hi, lo = max(l1, l2), min(l1, l2)
    return (hi + 0.05) / (lo + 0.05)


def theme_contrast(theme_name: str):
    """Contrast facts for a color theme, or ``None`` if the theme is unknown.

    Body = palette ``text`` on ``panel_bg`` (the main reading surface).
    Header = palette ``white`` on ``primary`` (panel header bars, large text).
    """
    from app.ppt_renderer import PALETTES

    pal = PALETTES.get(theme_name)
    if pal is None:
        return None
    body = wcag_ratio(tuple(pal.text), tuple(pal.panel_bg))
    header = wcag_ratio(tuple(pal.white), tuple(pal.primary))
    return {
        "theme": theme_name,
        "body_ratio": round(body, 2),
        "header_ratio": round(header, 2),
        "ok": body >= AA_NORMAL and header >= AA_LARGE,
    }


def contrast_ok(theme_name: str):
    """True if the theme passes WCAG AA, ``None`` if the theme is unknown."""
    tc = theme_contrast(theme_name)
    return None if tc is None else tc["ok"]
